from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from cash_tower.source_profile import (
    FILENAME_RULES,
    build_profile_report,
    distinct_worksheet_structures,
    missing_required_fields,
    profile_source_files,
    propose_mappings,
    save_unapproved_mapping_draft,
    snappay_header_consistency,
)


def _workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_filename_rules_identify_august_pilot_source_groups(tmp_path):
    for prefix, _ in FILENAME_RULES:
        _workbook(tmp_path / f"{prefix} August 2026.xlsx", ["A", "B"], [[1, 2]])
    (tmp_path / "unrelated.xlsx").write_text("not a workbook", encoding="utf-8")

    before = {
        path.name: (path.stat().st_size, path.stat().st_mtime_ns)
        for path in tmp_path.glob("*.xlsx")
        if path.name != "unrelated.xlsx"
    }
    profiles = profile_source_files(tmp_path)

    assert {profile["source_type"] for profile in profiles} == {
        prefix for prefix, _ in FILENAME_RULES
    }
    assert len(profiles) == 5
    assert all(profile["worksheets"][0]["row_count"] == 1 for profile in profiles)
    assert all(profile["worksheets"][0]["headers"] == ["A", "B"] for profile in profiles)
    assert before == {
        path.name: (path.stat().st_size, path.stat().st_mtime_ns)
        for path in tmp_path.glob("*.xlsx")
        if path.name != "unrelated.xlsx"
    }


def test_profile_reports_required_fields_as_unmapped_until_approved():
    profile = {
        "source_type": "SnapPay AR Transaction",
        "inferred_source": "SnapPay",
        "worksheets": [{"worksheet": "Transactions", "headers": ["Txn Ref"]}],
    }
    assert set(missing_required_fields(profile)) == {
        "Unmapped: account_type",
        "Unmapped: amount",
        "Unmapped: batch_number",
        "Unmapped: transaction_date",
        "Unmapped: transaction_id",
    }


def test_profile_compares_approved_mapping_to_observed_workbook_headers():
    profile = {
        "source_type": "SnapPay AR Transaction",
        "inferred_source": "SnapPay",
        "worksheets": [{
            "worksheet": "Transactions",
            "headers": ["Txn Ref", "Payment Type", "Batch", "Net", "Date"],
        }],
    }
    mappings = {
        "sources": {
            "SnapPay": {
                "worksheets": ["Transactions"],
                "required": ["transaction_id", "account_type", "batch_number", "amount", "transaction_date"],
                "columns": {
                    "transaction_id": ["Txn Ref"],
                    "account_type": ["Payment Type"],
                    "batch_number": ["Batch"],
                    "amount": ["Net"],
                    "transaction_date": ["Wrong Date Header"],
                },
            }
        },
        "local_ingestion": {
            "file_types": {
                "SnapPay AR Transaction*.xlsx": {"source": "SnapPay"}
            }
        },
    }
    assert missing_required_fields(profile, mappings) == ["transaction_date"]


def test_profile_accepts_derived_bmo_amount_mapping():
    profile = {
        "source_type": "BMO BLUEPAY",
        "inferred_source": "BMO",
        "worksheets": [{
            "worksheet": "Summary Report",
            "headers": ["Date", "Debit", "Credit", "Customer Reference"],
        }],
    }
    mappings = {
        "sources": {
            "BMO": {
                "worksheets": ["Summary Report"],
                "required": ["amount", "transaction_date"],
                "required_any": ["customer_reference"],
                "columns": {
                    "transaction_date": ["Date"],
                    "customer_reference": ["Customer Reference"],
                },
                "amount_calculation": {
                    "operation": "credit_minus_debit",
                    "credit": "Credit",
                    "debit": "Debit",
                },
            }
        },
        "local_ingestion": {
            "file_types": {"BMO BLUEPAY": {"source": "BMO"}}
        },
    }
    assert missing_required_fields(profile, mappings) == []


def test_profile_report_is_generated_in_memory_with_file_and_header_sheets(tmp_path):
    source = tmp_path / "SnapPay AR Transaction August 2026.xlsx"
    _workbook(source, ["Txn Ref", "Net"], [["TX-1", 12]])
    profiles = profile_source_files(tmp_path)

    content = build_profile_report(profiles)
    report_path = tmp_path / "profile.xlsx"
    report_path.write_bytes(content)
    workbook = load_workbook(report_path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Source Files", "Worksheets", "Headers", "Missing Required",
            "Distinct Structures", "SnapPay Consistency", "Identifier Risks",
            "Unapproved Proposals",
        ]
        assert workbook["Source Files"].max_row == 2
        assert workbook["Headers"].max_row == 3
        assert workbook["Missing Required"].max_row > 1
        assert workbook["Unapproved Proposals"].max_row > 1
        assert all(
            row[8] is False
            for row in workbook["Unapproved Proposals"].iter_rows(
                min_row=2, values_only=True
            )
        )
    finally:
        workbook.close()


def test_profile_requires_existing_local_directory(tmp_path):
    missing = tmp_path / "not-present"
    try:
        profile_source_files(missing)
    except FileNotFoundError as error:
        assert "does not exist" in str(error)
    else:
        raise AssertionError("A missing local folder must fail explicitly.")


def test_header_detection_skips_introductory_rows_and_counts_rows_after_header(tmp_path):
    path = tmp_path / "SnapPay AR Transaction daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Date"
    sheet.append(["August 2026 payment export"])
    sheet.append([])
    sheet.append(["Generated for Finance"])
    sheet.append(["PG Transaction ID", "Account Type", "Batch Number", "Paid Amount", "Payment Date"])
    sheet.append(["SYNTHETIC-ID", "CARD", "00042", 1.25, "2026-08-01"])
    workbook.save(path)

    profile = profile_source_files(tmp_path)[0]
    worksheet = profile["worksheets"][0]
    assert worksheet["header_row"] == 4
    assert worksheet["row_count"] == 1
    assert worksheet["headers"] == [
        "PG Transaction ID", "Account Type", "Batch Number", "Paid Amount", "Payment Date"
    ]
    assert "SYNTHETIC-ID" not in str(profile)


def test_snap_pay_daily_consistency_reports_header_row_variation(tmp_path):
    headers = ["PG Transaction ID", "Account Type", "Batch Number", "Paid Amount", "Payment Date"]
    first = Workbook()
    first_sheet = first.active
    first_sheet.title = "Date"
    first_sheet.append(["Synthetic introduction"])
    first_sheet.append([])
    first_sheet.append(["Generated for tests"])
    first_sheet.append(headers)
    first_sheet.append(["SYN-1", "CARD", "001", 1, "2026-08-01"])
    first.save(tmp_path / "SnapPay AR Transaction 20260801.xlsx")
    second = tmp_path / "SnapPay AR Transaction 20260802.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Date"
    sheet.append(headers)
    sheet.append(["SYN-2", "CARD", "002", 2, "2026-08-02"])
    workbook.save(second)

    profiles = profile_source_files(tmp_path)
    consistency = snappay_header_consistency(profiles)
    assert len(consistency) == 1
    assert consistency[0]["headers_consistent"]
    assert not consistency[0]["placement_consistent"]
    assert consistency[0]["header_rows"] == [1, 4]
    assert consistency[0]["distinct_structures"] == 2
    assert len(distinct_worksheet_structures(profiles)) == 2


def test_identifier_risk_counts_storage_without_exposing_values(tmp_path):
    path = tmp_path / "G.L. 1.1070 synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CashTransactionsGL"
    sheet.append(["Batch Num", "Document Num", "GL Date"])
    sheet.append([123, 456, "2026-08-01"])
    sheet["A2"].number_format = "General"
    sheet["B2"].number_format = "000000"
    workbook.save(path)

    profile = profile_source_files(tmp_path)[0]
    risks = {risk["header"]: risk for risk in profile["worksheets"][0]["identifier_risks"]}
    assert risks["Batch Num"]["numeric_cells"] == 1
    assert risks["Batch Num"]["zero_padded_format_cells"] == 0
    assert risks["Document Num"]["numeric_cells"] == 1
    assert risks["Document Num"]["zero_padded_format_cells"] == 1
    assert "456" not in str(profile)


def test_mapping_proposals_are_unapproved_and_rc_remains_separate(tmp_path):
    _workbook(
        tmp_path / "SnapPay AR Transaction synthetic.xlsx",
        ["PG Transaction ID", "Account Type", "Batch Number", "Paid Amount", "Payment Date"],
        [["SYNTHETIC-ID", "CARD", "1", 10, "2026-08-01"]],
    )
    _workbook(
        tmp_path / "RC BLUEPAY GENERAL synthetic.xlsx",
        ["Document Type", "Doc Number", "Batch Number", "G/L Date", "LT 1 Amount", "Reference 1"],
        [["synthetic", "1", "2", "2026-08-01", 10, "synthetic"]],
    )
    profiles = profile_source_files(tmp_path)
    proposals = propose_mappings(profiles)
    assert proposals
    assert all(proposal["approved"] is False for proposal in proposals)
    assert any(
        proposal["canonical_field"] == "transaction_id"
        and proposal["observed_header"] == "PG Transaction ID"
        and proposal["confidence"] == "High"
        for proposal in proposals
    )
    rc = next(profile for profile in profiles if profile["source_type"] == "RC BLUEPAY GENERAL")
    assert rc["inferred_source"] == "Unassigned (RC adapter decision required)"
    assert all(proposal["approved"] is False for proposal in proposals)

    draft_path = tmp_path / "data" / "proposals.yaml"
    save_unapproved_mapping_draft(proposals, draft_path)
    import yaml

    draft = yaml.safe_load(draft_path.read_text(encoding="utf-8"))
    assert draft["status"] == "draft_unapproved"
    assert draft["approved"] is False
    assert all(item["approved"] is False for item in draft["mappings"])
