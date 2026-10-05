from __future__ import annotations

import re
from io import BytesIO
from itertools import chain, islice
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

FILENAME_RULES = (
    ("SnapPay AR Transaction", "SnapPay"),
    ("BMO BLUEPAY", "BMO"),
    ("BLUEPAY Fiserv Bankcard", "BluePay"),
    ("RC BLUEPAY GENERAL", "Unassigned (RC adapter decision required)"),
    ("G.L. 1.1070", "JDE"),
)

REQUIRED_FIELDS = {
    "SnapPay": ["transaction_id", "account_type", "batch_number", "amount", "transaction_date"],
    "BluePay": ["transaction_id", "backend_id", "amount", "transaction_date"],
    "BMO": ["amount", "transaction_date"],
    "JDE": ["batch_number", "amount", "transaction_date"],
}
OPTIONAL_FIELDS = [
    "transaction_id", "account_type", "batch_number", "backend_id",
    "customer_reference", "amount", "transaction_date", "settlement_date", "status",
    "payment_type", "document_number", "document_type", "reference_1",
    "reference_2", "reference_3",
]
HEADER_SCAN_ROWS = 50
IDENTIFIER_HEADERS = {
    "id", "transaction id", "pg transaction id", "payment id", "backend id", "document number",
    "document num", "doc number", "batch number", "batch num",
    "customer reference", "bank reference",
}

MAPPING_CANDIDATES: dict[str, list[tuple[str, tuple[str, ...], str, str]]] = {
    "SnapPay AR Transaction": [
        ("transaction_id", ("PG Transaction ID",), "High", "Header names the payment-gateway transaction ID; verify identifiers were stored as text."),
        ("account_type", ("Account Type",), "High", "Exact canonical concept in the observed header."),
        ("batch_number", ("Batch Number",), "High", "Exact concept; numeric cell storage creates a leading-zero risk."),
        ("amount", ("Paid Amount",), "Medium", "Likely transaction amount, but signed/reversal and discount treatment need confirmation."),
        ("transaction_date", ("Payment Date",), "Medium", "Likely payment event date; confirm it is the intended transaction date."),
    ],
    "BLUEPAY Fiserv Bankcard": [
        ("transaction_id", ("id",), "Medium", "Likely BluePay transaction identifier, but the generic header 'id' needs source-owner confirmation."),
        ("backend_id", ("backend_id",), "High", "Exact observed identifier header; check numeric storage for leading-zero loss."),
        ("amount", ("amount",), "Medium", "Exact amount header; confirm sign convention and whether refunds/reversals are signed."),
        ("transaction_date", ("issue_date",), "Medium", "Candidate event date; 'settle_date' is a separate settlement-date candidate."),
        ("payment_type", ("payment_type",), "High", "Exact canonical concept in the observed header."),
        ("settlement_date", ("settle_date",), "High", "Exact settlement-date concept; distinct from transaction date."),
    ],
    "BMO BLUEPAY": [
        ("customer_reference", ("Customer Reference",), "High", "Exact observed reference label; confirm it contains the processor reference used for matching."),
        ("customer_reference", ("Bank Reference",), "Low", "Alternative bank-assigned reference; meaning differs from Customer Reference and requires review."),
        ("amount", ("Debit", "Credit"), "Low", "Two observed columns; a signed amount requires an approved debit/credit rule, not a direct single-column mapping."),
        ("settlement_date", ("Availability",), "Low", "Candidate funds-availability date; confirm whether this is the settlement date."),
        ("transaction_date", ("Date",), "Medium", "Observed transaction/posting date, not necessarily the settlement date."),
        ("payment_type", ("Transaction Description",), "Low", "Description may identify payment type but is not a dedicated type field."),
    ],
    "G.L. 1.1070": [
        ("batch_number", ("Batch Num",), "High", "Abbreviated batch header; numeric storage can remove leading zeroes."),
        ("amount", ("August",), "Low", "Period-specific header; confirm which ledger amount and sign convention it represents."),
        ("transaction_date", ("GL Date",), "High", "Exact general-ledger date concept."),
        ("document_number", ("Document Num",), "High", "Abbreviated document-number header; numeric storage can remove leading zeroes."),
        ("document_type", ("Document Type and Desc",), "Low", "Combined type/description header; a distinct document-type field is not evident."),
    ],
    "RC BLUEPAY GENERAL": [
        ("batch_number", ("Batch Number",), "High", "Exact batch concept; numeric storage creates a leading-zero risk."),
        ("amount", ("LT 1 Amount", "LT 1 Debit", "LT 1 Credit", "LT 2 Amount", "LT 2 Debit", "LT 2 Credit"), "Low", "Multiple ledger/ debit/credit amount columns; ledger and sign rule require Finance confirmation."),
        ("transaction_date", ("G/L Date",), "High", "Exact general-ledger date concept."),
        ("document_number", ("Doc Number",), "High", "Exact document-number concept; numeric storage creates a leading-zero risk."),
        ("document_type", ("Document Type", "Do Ty"), "Low", "Two plausible fields, one descriptive and one abbreviated; confirm intended canonical value."),
        ("reference_1", ("Reference 1",), "Low", "Reference field is present but its business meaning is unspecified."),
        ("reference_2", ("Reference 2",), "Low", "Reference field is present but its business meaning is unspecified."),
        ("reference_3", ("Reference 3",), "Low", "Reference field is present but its business meaning is unspecified."),
    ],
}


def detect_source_type(filename: str) -> tuple[str, str] | None:
    """Classify only the explicitly approved filename prefixes; never infer headers."""
    stem = Path(filename).stem.casefold()
    for prefix, source in FILENAME_RULES:
        if stem.startswith(prefix.casefold()):
            return prefix, source
    return None


def profile_source_files(root: str | Path) -> list[dict[str, Any]]:
    """Inspect workbook structure and identifier storage without changing source files."""
    root_path = Path(root).expanduser()
    if not root_path.exists():
        raise FileNotFoundError(f"Local synchronized SharePoint folder does not exist: {root_path}")
    if not root_path.is_dir():
        raise NotADirectoryError(f"Local synchronized SharePoint path is not a folder: {root_path}")

    profiles = []
    for path in sorted(root_path.rglob("*"), key=lambda item: str(item).casefold()):
        if not path.is_file() or path.suffix.casefold() != ".xlsx":
            continue
        detected = detect_source_type(path.name)
        if detected is None:
            continue
        source_type, inferred_source = detected
        file_profile: dict[str, Any] = {
            "source_type": source_type,
            "inferred_source": inferred_source,
            "inference_basis": f"Filename starts with {source_type!r}",
            "filename": path.name,
            "path": str(path.resolve()),
            "size_bytes": path.stat().st_size,
            "worksheets": [],
            "profile_error": None,
        }
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for worksheet in workbook.worksheets:
                    rows = iter(worksheet.iter_rows())
                    header_row, header_cells, buffered_rows = _detect_header(rows)
                    headers = [_header_text(cell.value) for cell in header_cells]
                    indexed_headers = [
                        (index, header)
                        for index, header in enumerate(headers)
                        if header
                    ]
                    identifier_columns = [
                        (index, header) for index, header in indexed_headers
                        if _normalized_header(header) in IDENTIFIER_HEADERS
                    ]
                    identifier_counts = {
                        header: {"numeric_cells": 0, "zero_padded_format_cells": 0, "text_leading_zero_cells": 0}
                        for _, header in identifier_columns
                    }
                    data_rows = 0
                    if header_row is not None:
                        for row in chain(buffered_rows, rows):
                            if _row_has_data(row):
                                data_rows += 1
                                _count_identifier_storage(row, identifier_columns, identifier_counts)
                    identifier_risks = [
                        {
                            "header": header,
                            **counts,
                            "numeric_storage_risk": counts["numeric_cells"] > 0,
                            "leading_zero_text_observed": counts["text_leading_zero_cells"] > 0,
                        }
                        for header, counts in identifier_counts.items()
                    ]
                    file_profile["worksheets"].append({
                        "worksheet": worksheet.title,
                        "row_count": data_rows,
                        "headers": [header for _, header in indexed_headers],
                        "header_row": header_row,
                        "identifier_risks": identifier_risks,
                    })
            finally:
                workbook.close()
        except (OSError, ValueError, KeyError) as exc:
            file_profile["profile_error"] = type(exc).__name__
        profiles.append(file_profile)
    return profiles


def distinct_worksheet_structures(profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Collapse repeated file sheets while retaining file counts and header placement."""
    structures: dict[tuple[Any, ...], dict[str, Any]] = {}
    for profile in profiles:
        for sheet in profile["worksheets"]:
            key = (
                profile["source_type"],
                sheet["worksheet"],
                sheet["header_row"],
                tuple(sheet["headers"]),
            )
            structure = structures.setdefault(key, {
                "source_type": profile["source_type"],
                "inferred_source": profile["inferred_source"],
                "worksheet": sheet["worksheet"],
                "header_row": sheet["header_row"],
                "headers": list(sheet["headers"]),
                "file_count": 0,
                "filenames": [],
                "row_count": 0,
                "identifier_risks": {},
            })
            structure["file_count"] += 1
            structure["filenames"].append(profile["filename"])
            structure["row_count"] += sheet["row_count"]
            for risk in sheet.get("identifier_risks", []):
                aggregate = structure["identifier_risks"].setdefault(risk["header"], {
                    "numeric_cells": 0,
                    "zero_padded_format_cells": 0,
                    "text_leading_zero_cells": 0,
                    "numeric_storage_risk": False,
                    "leading_zero_text_observed": False,
                })
                for field in ("numeric_cells", "zero_padded_format_cells", "text_leading_zero_cells"):
                    aggregate[field] += risk[field]
                aggregate["numeric_storage_risk"] |= risk["numeric_storage_risk"]
                aggregate["leading_zero_text_observed"] |= risk["leading_zero_text_observed"]
    return list(structures.values())


def propose_mappings(profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Generate explicitly unapproved mapping candidates from observed headers."""
    proposals = []
    for structure in distinct_worksheet_structures(profiles):
        observed = structure["headers"]
        for canonical, candidates, confidence, reason in MAPPING_CANDIDATES.get(
            structure["source_type"], []
        ):
            present = [header for header in candidates if header in observed]
            for header in present or ["<No exact observed header>"]:
                proposal_reason = reason
                if header == "<No exact observed header>":
                    proposal_reason += " No exact candidate header was observed; manual mapping is required."
                if header != "<No exact observed header>" and header in structure["identifier_risks"]:
                    risk = structure["identifier_risks"][header]
                    if risk["numeric_cells"]:
                        proposal_reason += (
                            f" {risk['numeric_cells']} identifier cells are stored numerically; "
                            "leading zeroes may already be lost."
                        )
                proposals.append({
                    "source_type": structure["source_type"],
                    "canonical_field": canonical,
                    "observed_header": header,
                    "worksheet": structure["worksheet"],
                    "confidence": confidence,
                    "reason": proposal_reason,
                    "header_row": structure["header_row"],
                    "profiled_file_count": structure["file_count"],
                    "approved": False,
                })
    return proposals


def snappay_header_consistency(profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Compare sheet names, exact headers, and header-row positions across daily files."""
    by_sheet: dict[str, dict[tuple[Any, ...], list[str]]] = {}
    for profile in profiles:
        if profile["source_type"] != "SnapPay AR Transaction":
            continue
        for sheet in profile["worksheets"]:
            signature = (sheet["header_row"], tuple(sheet["headers"]))
            by_sheet.setdefault(sheet["worksheet"], {}).setdefault(signature, []).append(
                profile["filename"]
            )
    results = []
    for worksheet, signatures in sorted(by_sheet.items()):
        results.append({
            "worksheet": worksheet,
            "files": sum(len(filenames) for filenames in signatures.values()),
            "distinct_structures": len(signatures),
            "header_rows": sorted({signature[0] for signature in signatures}),
            "headers_consistent": len({signature[1] for signature in signatures}) == 1,
            "placement_consistent": len({signature[0] for signature in signatures}) == 1,
            "files_by_header_row": {
                str(signature[0]): filenames
                for signature, filenames in signatures.items()
            },
        })
    return results


def save_unapproved_mapping_draft(
    proposals: list[dict[str, Any]], path: str | Path
) -> None:
    """Persist proposed mappings separately from the active mapping configuration."""
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    import yaml

    document = {
        "status": "draft_unapproved",
        "approved": False,
        "source_profile_only": True,
        "mappings": [
            {**proposal, "approved": False}
            for proposal in proposals
        ],
    }
    output_path.write_text(
        yaml.safe_dump(document, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )


def _detect_header(rows: Any) -> tuple[int | None, tuple[Any, ...], list[tuple[Any, ...]]]:
    buffered = list(islice(rows, HEADER_SCAN_ROWS))
    best_index = None
    best_score = 0
    for row_index, row in enumerate(buffered, start=1):
        text_cells = sum(
            isinstance(cell.value, str) and bool(cell.value.strip())
            for cell in row
        )
        if text_cells > best_score:
            best_index = row_index
            best_score = text_cells
    if best_index is None or best_score < 2:
        return None, (), buffered
    return best_index, buffered[best_index - 1], buffered[best_index:]


def _normalized_header(header: str) -> str:
    return " ".join(header.casefold().replace("_", " ").split())


def _row_has_data(row: tuple[Any, ...]) -> bool:
    return any(cell.value is not None and str(cell.value).strip() for cell in row)


def _count_identifier_storage(
    row: tuple[Any, ...],
    identifier_columns: list[tuple[int, str]],
    counts: dict[str, dict[str, int]],
) -> None:
    for index, header in identifier_columns:
        if index >= len(row):
            continue
        cell = row[index]
        value = cell.value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            counts[header]["numeric_cells"] += 1
            if re.match(r"^0{2,}(?:;|$)", cell.number_format):
                counts[header]["zero_padded_format_cells"] += 1
        elif isinstance(value, str) and value.startswith("0") and len(value) > 1:
            counts[header]["text_leading_zero_cells"] += 1


def missing_required_fields(
    profile: dict[str, Any], mappings: dict[str, Any] | None = None
) -> list[str]:
    """Compare approved canonical mappings to observed workbook headers."""
    source = profile["inferred_source"]
    file_types = mappings.get("local_ingestion", {}).get("file_types", {}) if mappings else {}
    file_type_config = file_types.get(
        profile["source_type"] + "*.xlsx",
        file_types.get(profile["source_type"], {}),
    )
    source = file_type_config.get("source", source)
    if source not in REQUIRED_FIELDS:
        return ["Source adapter not assigned"]
    if not mappings or source not in mappings.get("sources", {}):
        return [f"Unmapped: {field}" for field in REQUIRED_FIELDS[source]]
    source_mapping = mappings["sources"][source]
    configured_sheets = source_mapping.get("worksheets", [])
    matching_sheets = [
        sheet for sheet in profile["worksheets"]
        if not configured_sheets or sheet["worksheet"] in configured_sheets
    ]
    if not matching_sheets:
        return ["Configured worksheet not found"]

    columns = source_mapping.get("columns", {})
    absent = set()
    for sheet in matching_sheets:
        available = set(sheet["headers"])
        for field in REQUIRED_FIELDS[source]:
            if field == "amount" and source_mapping.get("amount_calculation"):
                continue
            candidates = columns.get(field, [])
            candidates = candidates if isinstance(candidates, list) else [candidates]
            if not candidates or not any(candidate in available for candidate in candidates):
                absent.add(field)
        required_any = source_mapping.get("required_any", [])
        if required_any and not any(
            any(candidate in available for candidate in (
                columns.get(field, []) if isinstance(columns.get(field, []), list)
                else [columns.get(field)]
            ) if candidate)
            for field in required_any
        ):
            absent.add("one of " + ", ".join(required_any))
    return sorted(absent)


def build_profile_report(
    profiles: list[dict[str, Any]], mappings: dict[str, Any] | None = None
) -> bytes:
    """Return a workbook metadata report as bytes; no report is saved to disk."""
    file_rows = []
    worksheet_rows = []
    header_rows = []
    for profile in profiles:
        file_rows.append({
            "inferred_source_type": profile["source_type"],
            "inferred_adapter": profile["inferred_source"],
            "filename": profile["filename"],
            "local_path": profile["path"],
            "size_bytes": profile["size_bytes"],
            "worksheet_names": "; ".join(sheet["worksheet"] for sheet in profile["worksheets"]),
            "data_rows": sum(sheet["row_count"] for sheet in profile["worksheets"]),
            "profile_error": profile["profile_error"] or "",
            "missing_required_fields": "; ".join(missing_required_fields(profile, mappings)),
        })
        for sheet in profile["worksheets"]:
            worksheet_rows.append({
                "filename": profile["filename"],
                "inferred_source_type": profile["source_type"],
                "worksheet": sheet["worksheet"],
                "data_rows": sheet["row_count"],
                "header_row": sheet["header_row"],
                "headers": " | ".join(sheet["headers"]),
            })
            for position, header in enumerate(sheet["headers"], start=1):
                header_rows.append({
                    "filename": profile["filename"],
                    "worksheet": sheet["worksheet"],
                    "column_number": position,
                    "header": header,
                })
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(file_rows, columns=[
            "inferred_source_type", "inferred_adapter", "filename", "local_path",
            "size_bytes", "worksheet_names", "data_rows", "missing_required_fields", "profile_error",
        ]).to_excel(writer, sheet_name="Source Files", index=False)
        pd.DataFrame(worksheet_rows, columns=[
            "filename", "inferred_source_type", "worksheet", "data_rows", "header_row", "headers",
        ]).to_excel(writer, sheet_name="Worksheets", index=False)
        pd.DataFrame(header_rows, columns=[
            "filename", "worksheet", "column_number", "header",
        ]).to_excel(writer, sheet_name="Headers", index=False)
        missing_rows = [
            {
                "filename": profile["filename"],
                "inferred_source_type": profile["source_type"],
                "inferred_adapter": profile["inferred_source"],
                "missing_required_field": field,
            }
            for profile in profiles
            for field in missing_required_fields(profile, mappings)
        ]
        pd.DataFrame(missing_rows, columns=[
            "filename", "inferred_source_type", "inferred_adapter", "missing_required_field",
        ]).to_excel(writer, sheet_name="Missing Required", index=False)
        structure_rows = [{
            "source_type": item["source_type"],
            "worksheet": item["worksheet"],
            "header_row": item["header_row"],
            "file_count": item["file_count"],
            "data_rows": item["row_count"],
            "headers": " | ".join(item["headers"]),
        } for item in distinct_worksheet_structures(profiles)]
        pd.DataFrame(structure_rows, columns=[
            "source_type", "worksheet", "header_row", "file_count", "data_rows", "headers",
        ]).to_excel(writer, sheet_name="Distinct Structures", index=False)
        consistency_rows = [{
            "worksheet": item["worksheet"],
            "files": item["files"],
            "distinct_structures": item["distinct_structures"],
            "header_rows": ", ".join(map(str, item["header_rows"])),
            "headers_consistent": item["headers_consistent"],
            "placement_consistent": item["placement_consistent"],
            "files_by_header_row": "; ".join(
                f"row {row}: {', '.join(filenames)}"
                for row, filenames in item["files_by_header_row"].items()
            ),
        } for item in snappay_header_consistency(profiles)]
        pd.DataFrame(consistency_rows, columns=[
            "worksheet", "files", "distinct_structures", "header_rows",
            "headers_consistent", "placement_consistent", "files_by_header_row",
        ]).to_excel(writer, sheet_name="SnapPay Consistency", index=False)
        risk_rows = [{
            "filename": profile["filename"],
            "source_type": profile["source_type"],
            "worksheet": sheet["worksheet"],
            "identifier_header": risk["header"],
            "numeric_cells": risk["numeric_cells"],
            "zero_padded_format_cells": risk["zero_padded_format_cells"],
            "text_leading_zero_cells": risk["text_leading_zero_cells"],
            "numeric_storage_risk": risk["numeric_storage_risk"],
            "leading_zero_text_observed": risk["leading_zero_text_observed"],
        } for profile in profiles
            for sheet in profile["worksheets"]
            for risk in sheet.get("identifier_risks", [])]
        pd.DataFrame(risk_rows, columns=[
            "filename", "source_type", "worksheet", "identifier_header",
            "numeric_cells", "zero_padded_format_cells", "text_leading_zero_cells",
            "numeric_storage_risk", "leading_zero_text_observed",
        ]).to_excel(writer, sheet_name="Identifier Risks", index=False)
        pd.DataFrame(propose_mappings(profiles)).to_excel(
            writer, sheet_name="Unapproved Proposals", index=False
        )
    return output.getvalue()


def _header_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text
